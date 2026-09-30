import csv
import io

from openpyxl import load_workbook

from skillgap.exports import HEADER, to_csv, to_xlsx
from skillgap.models import CandidateResult, Gap, OtherSkill, Recommendation, Skill


def sample(**overrides):
    base = dict(
        id="1", candidate="Maria", status="done", stage="done",
        skills=[Skill(id="fabric.lakehouse", name="Lakehouse", track="fabric", level=2, evidence="usou OneLake")],
        other_skills=[OtherSkill(name="Kubernetes", level=1, evidence="curso")],
        gaps=[Gap(skill="fabric.pipelines", name="Data Pipelines", track="fabric",
                  expected=2, current=0, severity="high")],
        recommendations=[Recommendation(course_id="c1", title="Fabric Pipelines", covers=["fabric.pipelines"],
                                        hours=8, link="https://example.com/c1")],
    )
    base.update(overrides)
    return CandidateResult(**base)


def parse(text):
    return list(csv.DictReader(io.StringIO(text)))


def test_csv_has_header_and_one_row_per_item():
    rows = parse(to_csv(sample()))
    assert list(rows[0].keys()) == HEADER
    assert [r["tipo"] for r in rows] == ["skill", "outra_skill", "gap", "treinamento"]
    gap = rows[2]
    assert (gap["nome"], gap["nivel_atual"], gap["nivel_esperado"], gap["severidade"]) == (
        "Data Pipelines", "0", "2", "high")
    assert rows[3]["horas"] == "8" and rows[3]["link"] == "https://example.com/c1"


def test_csv_neutralizes_formula_injection_from_cv_content():
    evil = sample(skills=[Skill(id="x", name="=HYPERLINK(\"http://evil\")", track="fabric",
                                level=1, evidence="+cmd|' /C calc'!A0")])
    row = parse(to_csv(evil))[0]
    assert row["nome"].startswith("'=") and row["evidencia"].startswith("'+")


def test_xlsx_has_expected_sheets_and_neutralizes_formulas():
    evil = sample(other_skills=[OtherSkill(name="@SUM(1+1)", level=1, evidence="-2+3")])
    workbook = load_workbook(io.BytesIO(to_xlsx(evil)))
    assert workbook.sheetnames == ["Skills", "Gaps", "Treinamentos"]
    values = [c.value for row in workbook["Skills"].iter_rows(min_row=2) for c in row]
    assert "'@SUM(1+1)" in values and "'-2+3" in values
    assert not any(isinstance(v, str) and v.startswith("=") for v in values)


def test_removes_illegal_control_characters_before_formula_check():
    """IMPORTANT: openpyxl rejects control chars like \x01, \x0b. Must remove them before prefixing."""
    # Skill name with control char \x01, evidence with \x0b and formula
    evil = sample(skills=[Skill(id="x", name="a\x01b", track="fabric", level=1, evidence="\x0b=1")])

    # to_xlsx should not raise IllegalCharacterError
    xlsx_bytes = to_xlsx(evil)
    workbook = load_workbook(io.BytesIO(xlsx_bytes))
    values = [c.value for row in workbook["Skills"].iter_rows(min_row=2) for c in row]

    # Control chars removed, but formula prefix applied to evidence
    assert "ab" in values  # control char removed
    assert "'=1" in values  # control char removed from evidence, formula prefixed

    # to_csv should show same cleaned values
    csv_text = to_csv(evil)
    rows = parse(csv_text)
    assert rows[0]["nome"] == "ab"  # control char removed
    assert rows[0]["evidencia"] == "'=1"  # control char removed and formula prefixed


def test_neutralizes_formulas_with_leading_whitespace():
    """MINOR: Formulas with leading whitespace (space, tab, invisible chars) must be caught and prefixed."""
    test_cases = [
        (" =1+1", "' =1+1"),  # leading space
        ("\n=1", "'\n=1"),    # leading newline
        ("﻿=1", "'﻿=1"),      # leading zero-width no-break space
        ("\xa0=1", "'\xa0=1"), # leading non-breaking space
        ("  \t@x", "'  \t@x"), # leading spaces and tab, formula prefix
    ]

    for malicious_input, expected_output in test_cases:
        skill_result = sample(skills=[Skill(id="x", name=malicious_input, track="fabric", level=1, evidence="safe")])

        # CSV test
        csv_rows = parse(to_csv(skill_result))
        assert csv_rows[0]["nome"] == expected_output, f"CSV: expected {expected_output!r}, got {csv_rows[0]['nome']!r}"

        # XLSX test
        workbook = load_workbook(io.BytesIO(to_xlsx(skill_result)))
        values = [c.value for row in workbook["Skills"].iter_rows(min_row=2) for c in row]
        assert expected_output in values, f"XLSX: expected {expected_output!r} in {values}"


def test_ordinary_values_unchanged_and_sole_minus_prefixed():
    """Guard tests: ordinary values unchanged, but lone '-' or '+' ARE prefixed (accepted trade-off)."""
    test_cases = [
        ("Lakehouse", "Lakehouse"),  # no change
        ("150M+ reads", "150M+ reads"),  # + not at start
        ("Data Pipelines", "Data Pipelines"),  # no change
        ("-", "'-"),  # lone minus is prefixed
        ("+", "'+"),  # lone plus is prefixed
    ]

    for input_val, expected_output in test_cases:
        skill_result = sample(skills=[Skill(id="x", name=input_val, track="fabric", level=1, evidence="safe")])

        # CSV test
        csv_rows = parse(to_csv(skill_result))
        assert csv_rows[0]["nome"] == expected_output, f"CSV: expected {expected_output!r}, got {csv_rows[0]['nome']!r}"

        # XLSX test
        workbook = load_workbook(io.BytesIO(to_xlsx(skill_result)))
        values = [c.value for row in workbook["Skills"].iter_rows(min_row=2) for c in row]
        assert expected_output in values, f"XLSX: expected {expected_output!r} in {values}"


def test_exports_write_empty_cell_for_unknown_hours():
    rec = Recommendation(course_id="c9", title="Sem horas", covers=["x"], hours=None, link="")
    result = sample(recommendations=[rec])
    row = parse(to_csv(result))[-1]
    assert row["tipo"] == "treinamento" and row["horas"] == "" and row["link"] == ""
    ws = load_workbook(io.BytesIO(to_xlsx(result)))["Treinamentos"]
    values = [c.value for c in ws[2]]
    assert values[HEADER.index("horas")] in (None, "")


def test_old_stored_result_json_without_new_recommendation_fields_still_loads():
    old = ('{"id":"1","status":"done","recommendations":[{"course_id":"c1","title":"T",'
           '"covers":["a"],"hours":8,"link":"https://x"}]}')
    [rec] = CandidateResult.model_validate_json(old).recommendations
    assert rec.hours == 8 and rec.provider is None and rec.kind is None and rec.verified is None
